"""源程序文档生成器。

软著申报对源程序文档的格式要求：

    - 前 30 页 + 后 30 页，合计 60 页（不足 60 页则全部提交）
    - 每页约 55 行
    - 页眉写明软件名称、版本号、页码

本模块按这些要求分页取材，并给每一行加上连续行号。
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

from . import docx_helpers as H
from .config import ProjectConfig

SEPARATOR_WIDTH = 80
CODE_FONT_SIZE = Pt(7.5)
CODE_LINE_SPACING = Pt(10.5)


def collect_lines(cfg: ProjectConfig) -> tuple[list[str], int]:
    """读取所有源文件，返回（行列表, 文件个数）。

    每个文件前插入一段分隔标记，并展开 Tab，避免 Word 里缩进错乱。
    """
    files = cfg.source_files()
    lines: list[str] = []

    for path in files:
        rel = _display_path(path, cfg.root)
        lines.append("# " + "=" * 68)
        lines.append(f"# 文件：{rel}")
        lines.append("# " + "=" * 68)
        lines.append("")

        try:
            content = path.read_text(encoding="utf-8")
        except UnicodeDecodeError:
            try:
                content = path.read_text(encoding="gbk")
            except Exception as exc:  # noqa: BLE001
                lines.append(f"# [错误] 无法解码该文件：{exc}")
                lines.append("")
                continue
        except Exception as exc:  # noqa: BLE001
            lines.append(f"# [错误] 无法读取该文件：{exc}")
            lines.append("")
            continue

        lines.extend(line.expandtabs(4) for line in content.split("\n"))
        lines.append("")
        lines.append("")

    return lines, len(files)


def paginate(lines: list[str], lines_per_page: int) -> list[list[str]]:
    """按每页行数切分。"""
    return [lines[i:i + lines_per_page]
            for i in range(0, len(lines), lines_per_page)]


def select_pages(pages: list[list[str]], first: int, last: int):
    """取前 first 页 + 后 last 页。

    返回 [(part, 全局页码, 该页的行), ...]，part 为 'front' 或 'back'。
    不足 first + last 页时全部返回。
    """
    total = len(pages)
    front = min(first, total)
    back = min(last, total - front) if total > front else 0

    selected = [("front", i + 1, pages[i]) for i in range(front)]
    back_start = total - back
    selected.extend(("back", back_start + i + 1, pages[back_start + i])
                    for i in range(back))
    return selected


def generate(cfg: ProjectConfig) -> Path:
    """生成源程序文档，返回输出文件路径。"""
    print("\n[1/4] 生成源程序文档 ...")

    lines, file_count = collect_lines(cfg)
    total_lines = len(lines)
    if total_lines == 0:
        raise SystemExit(
            "  [错误] 没有收集到任何源代码。\n"
            "         请检查 project.json 里的 source.dirs 和 source.include。"
        )

    pages = paginate(lines, cfg.source.lines_per_page)
    selected = select_pages(pages, cfg.source.first_pages, cfg.source.last_pages)

    print(f"      源文件 {file_count} 个，共 {total_lines} 行，"
          f"折合 {len(pages)} 页")
    print(f"      取前 {min(cfg.source.first_pages, len(pages))} 页 + "
          f"后 {len(selected) - min(cfg.source.first_pages, len(pages))} 页"
          f" = {len(selected)} 页")

    doc = Document()
    H.setup_page(doc, left_cm=2.0, right_cm=1.5, top_cm=2.0, bottom_cm=1.5,
                 base_font=H.SONG, base_size=Pt(9), line_spacing=Pt(14))

    H.add_cover(
        doc,
        big_title="源 程 序 文 档",
        subtitle="",
        info_lines=[
            f"软件名称：{cfg.name}",
            f"版本号：{cfg.version}",
            f"开发语言：{cfg.language}",
            f"代码规模：约 {total_lines} 行",
            f"文档页数：{len(selected)} 页",
        ],
        note="（源代码正文见后续页面）",
        rule_char="━",
        rule_count=40,
    )

    for part, page_no, page_lines in selected:
        doc.add_page_break()

        part_label = "（前部）" if part == "front" else "（后部）"
        H.add_paragraph(
            doc,
            f"{cfg.name} {cfg.version}  源程序文档 {part_label}  第 {page_no} 页",
            font_name=H.SONG, font_size=Pt(8.5), bold=True,
            alignment=WD_ALIGN_PARAGRAPH.CENTER, space_after=Pt(6),
        )
        H.add_separator(doc, char="─", count=SEPARATOR_WIDTH)

        for idx, line in enumerate(page_lines):
            global_no = (page_no - 1) * cfg.source.lines_per_page + idx + 1
            p = doc.add_paragraph()
            pf = p.paragraph_format
            pf.space_before = Pt(0)
            pf.space_after = Pt(0)
            pf.line_spacing = CODE_LINE_SPACING
            H.add_run(p, f"{global_no:5d}  {line}", H.MONO, CODE_FONT_SIZE)

        footer = doc.add_paragraph()
        footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        footer.paragraph_format.space_before = Pt(8)
        H.add_run(footer, f"— {page_no} —", H.SONG, Pt(8),
                  color=RGBColor(0xAA, 0xAA, 0xAA))

    out = cfg.output_path() / "源程序文档.docx"
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    print(f"      [OK] {out}")
    return out


def _display_path(path: Path, root: Path) -> str:
    try:
        return path.relative_to(root).as_posix()
    except ValueError:
        return path.as_posix()
