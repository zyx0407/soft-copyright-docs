"""用户操作手册生成器。

面向「不会用的人照着做」，所以每个功能模块都带编号步骤和截图占位框。
生成的 Word 文档打开后，往占位框里贴截图即可，不用自己排版。
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

from . import docx_helpers as H
from .config import ProjectConfig


def generate(cfg: ProjectConfig) -> Path:
    """生成用户操作手册，返回输出文件路径。"""
    print("\n[3/4] 生成用户操作手册 ...")

    doc = Document()
    H.setup_page(doc, left_cm=2.2, right_cm=2.2, top_cm=1.8, bottom_cm=1.5)
    H.setup_headings(doc)

    info_lines = [f"版本号：{cfg.version}"]
    if cfg.name_en:
        info_lines.insert(0, f"英文名称：{cfg.name_en}")
    if cfg.developer:
        info_lines.append(f"开发单位：{cfg.developer}")
    info_lines.append("")
    info_lines.append("（本文档为软件著作权登记材料）")

    H.add_cover(doc, big_title="用户操作手册", subtitle="",
                info_lines=info_lines)

    H.add_toc_field(doc)

    # ── 第一章 软件概述 ──────────────────────────
    H.add_page_break(doc)
    H.add_heading_styled(doc, "第一章  软件概述", level=1)

    H.add_heading_styled(doc, "1.1  软件简介", level=2)
    if cfg.summary:
        H.add_indented_paragraph(doc, cfg.summary)
    else:
        H.add_indented_paragraph(
            doc, f"{cfg.name}（以下简称“本软件”）当前版本 {cfg.version}。")

    if cfg.features:
        H.add_heading_styled(doc, "1.2  核心功能", level=2)
        for f in cfg.features:
            text = f"**{f.title}**：{f.detail}" if f.detail else f"**{f.title}**"
            H.add_bullet(doc, text, font_size=Pt(10.5))

    # ── 第二章 安装与启动 ────────────────────────
    H.add_page_break(doc)
    H.add_heading_styled(doc, "第二章  安装与启动", level=1)

    H.add_heading_styled(doc, "2.1  系统要求", level=2)

    if cfg.hardware:
        H.add_paragraph(doc, "硬件环境：", font_name=H.HEI, font_size=Pt(11),
                        bold=True, space_after=Pt(4))
        for item in cfg.hardware:
            H.add_bullet(doc, item, font_size=Pt(10.5))

    if cfg.software:
        H.add_paragraph(doc, "软件环境：", font_name=H.HEI, font_size=Pt(11),
                        bold=True, space_after=Pt(4))
        for item in cfg.software:
            H.add_bullet(doc, item, font_size=Pt(10.5))

    H.add_heading_styled(doc, "2.2  启动软件", level=2)
    H.add_numbered_step(doc, 1, "确认计算机满足 2.1 节所列的系统要求。")
    H.add_numbered_step(
        doc, 2, "将本软件的发布包完整解压到本地磁盘的任意目录下"
                "（建议路径中不要包含中文或空格）。")
    H.add_numbered_step(doc, 3, f"进入解压后的目录，双击主程序即可启动 {cfg.name}。")
    H.add_numbered_step(doc, 4, "首次启动时，软件会自动完成初始化，"
                                "稍候片刻即可进入主界面。")

    H.add_note_after_space(
        doc,
        "请勿将软件放在 U 盘或网络盘中直接运行，也不要删除同目录下的"
        "依赖文件，否则可能导致启动失败。",
        "注意")

    # ── 第三章 功能操作说明 ──────────────────────
    H.add_page_break(doc)
    H.add_heading_styled(doc, "第三章  功能操作说明", level=1)
    H.add_paragraph(
        doc,
        f"本章按功能模块逐个说明操作方法，共 {len(cfg.modules)} 个模块。"
        "每个模块下方的虚线框用于放置界面截图，可直接在 Word 中插入图片。",
        font_size=Pt(10.5), space_after=Pt(10))

    if not cfg.modules:
        H.add_note(
            doc,
            "尚未在 project.json 的 modules 字段里描述功能模块，"
            "请补充后重新生成。",
            "重要")

    for i, m in enumerate(cfg.modules, start=1):
        H.add_page_break(doc)
        H.add_heading_styled(doc, f"3.{i}  {m.name}", level=2)

        if m.desc:
            H.add_indented_paragraph(doc, m.desc)

        if m.steps:
            H.add_paragraph(doc, "操作步骤：", font_name=H.HEI,
                            font_size=Pt(11), bold=True, space_after=Pt(6))
            for n, step in enumerate(m.steps, start=1):
                H.add_numbered_step(doc, n, step)

        H.add_paragraph(doc, "", font_size=Pt(6), space_after=Pt(0))
        H.add_paragraph(doc, "界面截图：", font_name=H.HEI, font_size=Pt(11),
                        bold=True, space_after=Pt(6))
        H.add_image_placeholder(doc, width_cm=15.0, height_cm=8.5)

    # ── 第四章 注意事项 ──────────────────────────
    H.add_page_break(doc)
    H.add_heading_styled(doc, "第四章  使用注意事项", level=1)

    H.add_heading_styled(doc, "4.1  数据安全", level=2)
    H.add_indented_paragraph(
        doc,
        "软件的处理结果会输出到用户指定的目录中，建议在重要操作前"
        "自行备份原始数据。软件不会修改、移动或删除用户的原始文件。")
    H.add_note_after_space(
        doc, "请定期备份原始数据文件，避免因磁盘故障造成不可恢复的丢失。",
        "警告")

    H.add_heading_styled(doc, "4.2  运行环境", level=2)
    H.add_indented_paragraph(
        doc,
        "软件对显示分辨率有一定要求。若界面显示不完整，"
        "请调高显示器分辨率或缩放比例后重新启动软件。")

    H.add_heading_styled(doc, "4.3  其他", level=2)
    H.add_indented_paragraph(
        doc,
        "若软件在使用过程中出现异常，请先关闭软件后重新启动；"
        "问题若仍然存在，请记录操作步骤与界面提示信息，"
        "连同日志文件一并反馈给开发人员。")

    out = cfg.output_path() / "用户操作手册.docx"
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    print(f"      [OK] {out}")
    return out
