"""软件界面截图文档生成器。

用于软著申报（尤其是被要求补正时）：一页一张截图，
配好模块名和功能说明，虚线框里贴图即可。
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

from . import docx_helpers as H
from .config import ProjectConfig


def generate(cfg: ProjectConfig) -> Path:
    """生成界面截图文档，返回输出文件路径。"""
    print("\n[4/4] 生成界面截图文档 ...")

    modules = cfg.modules
    doc = Document()
    H.setup_page(doc, left_cm=2.0, right_cm=2.0, top_cm=1.8, bottom_cm=1.5)
    H.setup_headings(doc)

    info_lines = [
        f"软件名称：{cfg.name}",
    ]
    if cfg.name_en:
        info_lines.append(f"英文名称：{cfg.name_en}")
    info_lines.append(f"版本号：{cfg.version}")
    info_lines.append("")
    info_lines.append(f"开发单位：{cfg.developer}" if cfg.developer else "开发单位：（请填写）")
    info_lines.append(
        f"完成日期：{cfg.completion_date}" if cfg.completion_date
        else "完成日期：（请填写）")

    H.add_cover(
        doc,
        big_title="软 件 界 面 截 图 文 档",
        subtitle="",
        info_lines=info_lines,
        note=f"（共 {len(modules)} 张截图）",
    )

    # ── 功能模块与截图索引 ───────────────────────
    H.add_page_break(doc)
    H.add_paragraph(doc, "软件功能模块与截图索引", font_name=H.HEI,
                    font_size=Pt(16), bold=True,
                    alignment=WD_ALIGN_PARAGRAPH.CENTER, space_after=Pt(16))

    if modules:
        H.add_grid_table(
            doc,
            headers=["序号", "功能模块", "说明"],
            rows=[[m.id or str(i), m.name, m.desc] for i, m in enumerate(modules, start=1)],
            widths_cm=[1.6, 4.4, 9.5],
        )
    else:
        H.add_note(
            doc,
            "尚未在 project.json 的 modules 字段里描述功能模块，"
            "请补充后重新生成。",
            "重要")

    # ── 每页一张截图 ─────────────────────────────
    for i, m in enumerate(modules, start=1):
        H.add_page_break(doc)

        H.add_paragraph(doc, f"{m.id or i}. {m.name}", font_name=H.HEI,
                        font_size=Pt(15), bold=True,
                        alignment=WD_ALIGN_PARAGRAPH.CENTER, space_after=Pt(10))

        if m.desc:
            H.add_paragraph(doc, m.desc, font_size=Pt(10.5),
                            alignment=WD_ALIGN_PARAGRAPH.CENTER,
                            space_after=Pt(8),
                            color=H.DARK)

        H.add_image_placeholder(doc, width_cm=16.0, height_cm=10.5)

        if m.steps:
            H.add_paragraph(doc, "", font_size=Pt(6), space_after=Pt(0))
            H.add_paragraph(doc, "操作说明：", font_name=H.HEI,
                            font_size=Pt(10.5), bold=True, space_after=Pt(4))
            for n, step in enumerate(m.steps, start=1):
                H.add_numbered_step(doc, n, step, font_size=Pt(10))

    out = cfg.output_path() / "界面截图文档.docx"
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    print(f"      [OK] {out}")
    return out
