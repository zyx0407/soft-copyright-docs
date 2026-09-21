"""软件说明书生成器。

说明书讲「这个软件是什么、能干什么、跑在什么环境上」，
操作手册讲「怎么一步步用它」——两者是软著申报里的两份不同材料。
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.shared import Cm, Pt

from . import docx_helpers as H
from .config import ProjectConfig


def generate(cfg: ProjectConfig) -> Path:
    """生成软件说明书，返回输出文件路径。"""
    print("\n[2/4] 生成软件说明书 ...")

    doc = Document()
    H.setup_page(doc, left_cm=2.2, right_cm=2.2, top_cm=1.8, bottom_cm=1.5)
    H.setup_headings(doc)

    info_lines = [f"版本号：{cfg.version}"]
    if cfg.name_en:
        info_lines.insert(0, f"英文名称：{cfg.name_en}")
    if cfg.developer:
        info_lines.append(f"开发单位：{cfg.developer}")
    if cfg.completion_date:
        info_lines.append(f"完成日期：{cfg.completion_date}")

    H.add_cover(doc, big_title="软 件 说 明 书", subtitle="",
                info_lines=info_lines)

    H.add_toc_field(doc)

    # ── 第一章 软件概述 ──────────────────────────
    H.add_page_break(doc)
    H.add_heading_styled(doc, "第一章  软件概述", level=1)

    H.add_heading_styled(doc, "1.1  软件简介", level=2)
    if cfg.summary:
        H.add_indented_paragraph(doc, cfg.summary)
    else:
        H.add_indented_paragraph(
            doc,
            f"{cfg.name}是一款{cfg.language}开发的软件。"
            "（此处可在 project.json 的 summary 字段里填写软件简介）"
        )
    if cfg.name_en:
        H.add_indented_paragraph(
            doc,
            f"本软件英文名称为 {cfg.name_en}，当前版本 {cfg.version}。"
        )

    if cfg.features:
        H.add_heading_styled(doc, "1.2  核心功能", level=2)
        for f in cfg.features:
            text = f"**{f.title}**：{f.detail}" if f.detail else f"**{f.title}**"
            H.add_bullet(doc, text, font_size=Pt(10.5))

    if cfg.glossary:
        H.add_heading_styled(doc, "1.3  术语说明", level=2)
        H.add_paragraph(
            doc, "本说明书中出现的术语与缩写含义如下：",
            font_size=Pt(10.5), space_after=Pt(8))
        H.add_grid_table(
            doc,
            headers=["术语", "全称", "说明"],
            rows=[[g.term, g.full, g.desc] for g in cfg.glossary],
            widths_cm=[2.6, 4.6, 8.3],
        )

    # ── 第二章 运行环境 ──────────────────────────
    H.add_page_break(doc)
    H.add_heading_styled(doc, "第二章  运行环境", level=1)

    if cfg.hardware:
        H.add_heading_styled(doc, "2.1  硬件环境", level=2)
        for item in cfg.hardware:
            H.add_bullet(doc, item, font_size=Pt(10.5))

    if cfg.software:
        H.add_heading_styled(doc, "2.2  软件环境", level=2)
        for item in cfg.software:
            H.add_bullet(doc, item, font_size=Pt(10.5))

    # ── 第三章 功能模块 ──────────────────────────
    if cfg.modules:
        H.add_page_break(doc)
        H.add_heading_styled(doc, "第三章  功能模块说明", level=1)
        H.add_paragraph(
            doc,
            f"本软件共包含 {len(cfg.modules)} 个功能模块，说明如下。",
            font_size=Pt(10.5), space_after=Pt(10))

        for i, m in enumerate(cfg.modules, start=1):
            H.add_heading_styled(doc, f"3.{i}  {m.name}", level=2)
            if m.desc:
                H.add_indented_paragraph(doc, m.desc)
            if m.steps:
                H.add_paragraph(doc, "操作流程：", font_name=H.HEI,
                                font_size=Pt(10.5), bold=True, space_after=Pt(4))
                for n, step in enumerate(m.steps, start=1):
                    H.add_numbered_step(doc, n, step, font_size=Pt(10.5))
            H.add_paragraph(doc, "", font_size=Pt(6), space_after=Pt(0))

        H.add_note(
            doc,
            f"完整的操作步骤与界面截图见配套的《{cfg.name} 用户操作手册》。",
            "提示")

    out = cfg.output_path() / "软件说明书.docx"
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    print(f"      [OK] {out}")
    return out
