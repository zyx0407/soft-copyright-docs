"""检查 `python -m ruanzhu` 生成的四份软著材料是否装上得了台面。

用法：
    python verify_materials.py                     # 检查当前目录下的 软著申请文档/
    python verify_materials.py 我的输出目录
    python verify_materials.py 我的输出目录 --name 软件名 --version V1.0

判据是机器能查的那一半：文件在不在、有没有内容、四份的软件名与版本号是否一致、
源程序文档的页码是否连续、每页标题与页脚对不对得上、代码行号是否从 1 连续编到末尾。
查不了的（目录域要手动更新、截图要人贴）会单独列成「需要人工」。

退出码：有 FAIL 为 1，只有 WARN 为 0。
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

try:
    from docx import Document
except ImportError:  # pragma: no cover
    print("[FAIL] 缺少 python-docx，先运行：pip install python-docx")
    raise SystemExit(1)


def _utf8_stdout() -> None:
    """Windows 下把 stdout/stderr 固定成 UTF-8，否则输出被当成 ANSI 解码会乱码。"""
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(encoding="utf-8")  # type: ignore[union-attr]
        except (AttributeError, ValueError):
            pass


_utf8_stdout()

MATERIALS = {
    "源程序文档.docx": "source",
    "软件说明书.docx": "spec",
    "用户操作手册.docx": "manual",
    "界面截图文档.docx": "screenshots",
}

PAGE_NUMBER_RE = re.compile(r"^—\s*(\d+)\s*—$")
PAGE_TITLE_RE = re.compile(r"源程序文档\s*（(前部|后部)）\s*第\s*(\d+)\s*页")
LINE_NO_RE = re.compile(r"^ *(\d{1,6}) {2}")
NAME_RE = re.compile(r"^软件名称[：:]\s*(.+)$")
VERSION_RE = re.compile(r"^版本号[：:]\s*(.+)$")


class Report:
    def __init__(self) -> None:
        self.fails: list[str] = []
        self.warns: list[str] = []
        self.ok: list[str] = []

    def fail(self, msg: str) -> None:
        self.fails.append(msg)

    def warn(self, msg: str) -> None:
        self.warns.append(msg)

    def pass_(self, msg: str) -> None:
        self.ok.append(msg)

    def dump(self) -> int:
        for m in self.ok:
            print(f"[OK]   {m}")
        for m in self.warns:
            print(f"[WARN] {m}")
        for m in self.fails:
            print(f"[FAIL] {m}")
        print()
        print(f"合计：PASS {len(self.ok)} · WARN {len(self.warns)} · FAIL {len(self.fails)}")
        if self.fails:
            print("有 FAIL：先把上面几项修掉，再谈提交。")
        else:
            print("机器能查的部分都过了；下面三件事做不了，得人工。")
        print()
        print("需要人工：")
        print("  1. 用 Word 打开四份文档，目录处右键 → 更新域 → 更新整个目录")
        print("  2. 界面截图文档里把虚线占位框换成真实截图")
        print("  3. 确认源代码著作权归属（合作开发先确认协议）")
        return 1 if self.fails else 0


def page_breaks(doc) -> int:
    """文档里的硬分页符个数。"""
    count = 0
    for p in doc.paragraphs:
        for r in p.runs:
            if 'w:type="page"' in r._element.xml:
                count += 1
    return count


def all_text(doc) -> str:
    """正文 + 表格里的全部文字（软著材料的封面信息常放在表格里）。"""
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def first_match(doc, pattern: re.Pattern[str]) -> str | None:
    for p in doc.paragraphs:
        m = pattern.match(p.text.strip())
        if m:
            return m.group(1).strip()
    return None


def check_source_doc(doc, rep: Report) -> None:
    """源程序文档：页码、页标题、行号三件事必须自洽。"""
    page_no_marks = [int(m.group(1)) for p in doc.paragraphs
                     if (m := PAGE_NUMBER_RE.match(p.text.strip()))]
    breaks = page_breaks(doc)

    if not page_no_marks:
        rep.fail("源程序文档：一个页脚页码都没找到，文档可能没生成完整")
        return

    expected = list(range(1, len(page_no_marks) + 1))
    if page_no_marks != expected:
        rep.fail(f"源程序文档：页脚页码不连续（读到 {page_no_marks[:3]}…{page_no_marks[-3:]}）")
    elif len(page_no_marks) != breaks:
        rep.fail(f"源程序文档：页脚页数 {len(page_no_marks)} 与分页符 {breaks} 对不上")
    else:
        rep.pass_(f"源程序文档：共 {len(page_no_marks)} 页，页码连续")

    titles = [(m.group(1), int(m.group(2))) for p in doc.paragraphs
              if (m := PAGE_TITLE_RE.search(p.text))]
    if len(titles) != len(page_no_marks):
        rep.fail(f"源程序文档：每页标题行应有 {len(page_no_marks)} 条，实际 {len(titles)} 条")
    else:
        mismatched = [i + 1 for i, (_, no) in enumerate(titles) if no != i + 1]
        if mismatched:
            rep.fail(f"源程序文档：页标题的页码与顺序对不上，第 {mismatched[:5]} 页")
        else:
            rep.pass_("源程序文档：每页标题行的页码与页脚一致")

    line_nos = []
    for p in doc.paragraphs:
        m = LINE_NO_RE.match(p.text)
        if m:
            line_nos.append(int(m.group(1)))
    if not line_nos:
        rep.fail("源程序文档：没找到连续行号，可能根本没取到源代码")
    elif line_nos[0] != 1:
        rep.fail(f"源程序文档：行号应从 1 开始，实际从 {line_nos[0]} 开始")
    elif line_nos != list(range(1, len(line_nos) + 1)):
        bad = next(i for i, v in enumerate(line_nos) if v != i + 1)
        rep.fail(f"源程序文档：行号在第 {bad + 1} 个代码行处断开（读到 {line_nos[bad]}）")
    else:
        rep.pass_(f"源程序文档：行号 1–{len(line_nos)} 连续，共 {len(line_nos)} 行代码")


def check_prose_doc(name: str, doc, rep: Report, want_toc: bool) -> None:
    text = all_text(doc)
    if len(text.strip()) < 200:
        rep.fail(f"{name}：内容几乎为空，别提交")
        return

    toc = sum(1 for p in doc.paragraphs for r in p.runs if "TOC" in r._element.xml)
    if want_toc and toc == 0:
        rep.warn(f"{name}：没找到目录域，封面/目录可能只剩提示文字")
    elif want_toc:
        rep.pass_(f"{name}：目录域已插入（打开后仍需手动更新一次）")

    placeholders = text.count("（请填写）")
    if placeholders:
        rep.warn(f"{name}：还有 {placeholders} 处「（请填写）」没填（开发单位/完成日期）")

    if "**" in text:
        rep.warn(f"{name}：正文出现 Markdown 残留 **，检查 project.json 里是不是混进了加粗标记")


def parse_args(argv: list[str]) -> tuple[Path, str | None, str | None]:
    """`[输出目录] [--name 值] [--version 值]`，键值也接受 --name=值。"""
    out_dir: Path | None = None
    name: str | None = None
    version: str | None = None
    i = 0
    while i < len(argv):
        arg = argv[i]
        key, _, inline = arg.partition("=")
        if key in ("--name", "--version"):
            value = inline or (argv[i + 1] if i + 1 < len(argv) else "")
            if not inline:
                i += 1
            if key == "--name":
                name = value or None
            else:
                version = value or None
        elif arg.startswith("--"):
            raise SystemExit(f"未知参数：{arg}")
        elif out_dir is None:
            out_dir = Path(arg)
        i += 1
    return out_dir or Path("软著申请文档"), name, version


def main(argv: list[str]) -> int:
    out_dir, expect_name, expect_version = parse_args(argv)

    rep = Report()
    if not out_dir.is_dir():
        rep.fail(f"输出目录不存在：{out_dir}")
        return rep.dump()

    docs = {}
    for fn in MATERIALS:
        path = out_dir / fn
        if not path.is_file():
            rep.fail(f"缺文件：{fn}")
            continue
        if path.stat().st_size < 5_000:
            rep.fail(f"{fn}：只有 {path.stat().st_size} 字节，疑似空文档")
            continue
        docs[fn] = Document(path)
        rep.pass_(f"{fn}：{path.stat().st_size // 1024} KB")

    if "源程序文档.docx" in docs:
        check_source_doc(docs["源程序文档.docx"], rep)
        name = first_match(docs["源程序文档.docx"], NAME_RE)
        version = first_match(docs["源程序文档.docx"], VERSION_RE)
    else:
        name = version = None

    if expect_name and name and name != expect_name:
        rep.fail(f"源程序文档里的软件名是「{name}」，与期望「{expect_name}」不一致")
    if expect_version and version and version != expect_version:
        rep.fail(f"版本号是「{version}」，与期望「{expect_version}」不一致")

    # 四份材料的软件名、版本号必须一致
    names, versions = set(), set()
    for fn, doc in docs.items():
        n = first_match(doc, NAME_RE)
        v = first_match(doc, VERSION_RE)
        if n:
            names.add(n)
        if v:
            versions.add(v)
        if name and n is None and name not in all_text(doc):
            rep.fail(f"{fn}：正文里找不到软件名「{name}」")
    if len(names) > 1:
        rep.fail(f"四份材料的软件名不一致：{sorted(names)}")
    elif names:
        rep.pass_(f"四份材料软件名一致：{sorted(names)[0]}")
    if len(versions) > 1:
        rep.fail(f"四份材料的版本号不一致：{sorted(versions)}")
    elif versions:
        rep.pass_(f"四份材料版本号一致：{sorted(versions)[0]}")

    for fn in ("软件说明书.docx", "用户操作手册.docx"):
        if fn in docs:
            check_prose_doc(fn, docs[fn], rep, want_toc=True)
    if "界面截图文档.docx" in docs:
        check_prose_doc("界面截图文档.docx", docs["界面截图文档.docx"], rep, want_toc=False)
        breaks = page_breaks(docs["界面截图文档.docx"])
        rep.pass_(f"界面截图文档：{breaks} 页占位（每页贴一张截图）")

    return rep.dump()


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
