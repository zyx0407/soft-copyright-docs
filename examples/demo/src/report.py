"""分析结果导出为 Word 报告。"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from stats import Summary


def build_report(title: str, sections: dict, out_path: str | Path) -> Path:
    """把若干分析板块写成一个 Word 文档。

    sections 的键是板块标题，值是字符串列表（每个元素一段）。
    """
    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = "宋体"
    style.font.size = Pt(11)

    heading = doc.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = heading.add_run(title)
    run.font.size = Pt(18)
    run.font.bold = True

    stamp = doc.add_paragraph()
    stamp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    stamp_run = stamp.add_run(
        f"生成时间：{datetime.now():%Y-%m-%d %H:%M:%S}")
    stamp_run.font.size = Pt(9)

    for name, paragraphs in sections.items():
        h = doc.add_heading(name, level=1)
        for r in h.runs:
            r.font.name = "黑体"
        for text in paragraphs:
            p = doc.add_paragraph()
            p.paragraph_format.first_line_indent = Pt(22)
            p.paragraph_format.line_spacing = Pt(20)
            p.add_run(text)

    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)
    return out_path


def summary_to_rows(summaries: list[Summary]) -> list[list[str]]:
    """把描述性统计结果整理成表格行，便于粘进 Word 或导出 Excel。"""
    rows = []
    for s in summaries:
        rows.append([
            s.name,
            str(s.count),
            f"{s.mean:.4f}",
            f"{s.std:.4f}",
            f"{s.median:.4f}",
            f"{s.q1:.4f}",
            f"{s.q3:.4f}",
            f"{s.cv:.4f}",
        ])
    return rows


HEADERS = ["变量", "样本数", "均值", "标准差", "中位数", "Q1", "Q3", "变异系数"]


def describe(summaries: list[Summary]) -> list[str]:
    """生成一段文字结论，说明各变量的离散程度。"""
    lines = []
    for s in summaries:
        if s.count == 0:
            lines.append(f"{s.name}：没有可用数据。")
            continue
        level = "较集中" if s.cv < 0.15 else ("中等离散" if s.cv < 0.3 else "离散较大")
        lines.append(
            f"{s.name}：共 {s.count} 个有效值，均值 {s.mean:.4f}，"
            f"标准差 {s.std:.4f}，变异系数 {s.cv:.4f}，数据{level}。"
        )
    return lines
